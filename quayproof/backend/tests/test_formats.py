import io
from app.config import Settings
from app.parsers import parse_document


def test_docx_tables():
    from docx import Document
    doc=Document();doc.add_paragraph('SHIPPING INSTRUCTION')
    row=doc.add_table(rows=1,cols=2).rows[0]
    row.cells[0].text='Gross Weight';row.cells[1].text='12,500 kg'
    out=io.BytesIO();doc.save(out)
    blocks=parse_document('si.docx',out.getvalue(),Settings())
    assert any('12,500 kg' in b['text'] and b['location'].get('table')==1 for b in blocks)


def test_xlsx_cell_evidence():
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Shipping'
    sheet.append(['SHIPPING INSTRUCTION']);sheet.append(['Gross Weight','12,500 kg'])
    out=io.BytesIO();book.save(out)
    blocks=parse_document('si.xlsx',out.getvalue(),Settings())
    assert blocks[1]['location']=={'sheet':'Shipping','cells':'A2, B2'}


def test_native_pdf_page_evidence():
    from reportlab.pdfgen.canvas import Canvas
    out=io.BytesIO();pdf=Canvas(out);pdf.drawString(60,750,'SHIPPING INSTRUCTION');pdf.drawString(60,720,'Gross Weight: 12,500 kg');pdf.save()
    blocks=parse_document('si.pdf',out.getvalue(),Settings())
    assert blocks[1]['location']=={'page':1} and blocks[1]['bbox'] and blocks[1]['method']=='native'


def test_scanned_pdf_ocr():
    import shutil
    import pytest
    from PIL import Image,ImageDraw,ImageFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    if not shutil.which('tesseract'):pytest.skip('Tesseract executable not installed')
    image=Image.new('RGB',(1200,500),'white');draw=ImageDraw.Draw(image)
    try:font=ImageFont.truetype('DejaVuSans.ttf',40)
    except OSError:font=ImageFont.load_default(size=40)
    draw.text((50,50),'SHIPPING INSTRUCTION',fill='black',font=font)
    draw.text((50,140),'Gross Weight: 12500 kg',fill='black',font=font)
    out=io.BytesIO();pdf=Canvas(out,pagesize=(600,250));pdf.drawImage(ImageReader(image),0,0,600,250);pdf.save()
    blocks=parse_document('scan.pdf',out.getvalue(),Settings())
    assert any('12500' in b['text'] and b['method']=='tesseract' and b['bbox'] for b in blocks)
