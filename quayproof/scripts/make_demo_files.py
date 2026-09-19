#!/usr/bin/env python3
"""Generate original synthetic files for upload demos; no participant data."""
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))
from app.demo import BASE

def main():
    folder=ROOT/'demo-files';folder.mkdir(exist_ok=True)
    samples={'instruction.txt':'SHIPPING INSTRUCTION\n'+BASE,'draft-matching.txt':'DRAFT BILL OF LADING\n'+BASE,'draft-mismatch.txt':'DRAFT BILL OF LADING\n'+BASE.replace('12,500 kg','12,800 kg'),'draft-unresolved.txt':'DRAFT BILL OF LADING\n'+BASE.replace('12,500 kg','TBC'),'draft-corrected.txt':'DRAFT BILL OF LADING\n'+BASE}
    for name,text in samples.items():(folder/name).write_text(text,encoding='utf-8')
    (folder/'unreadable.pdf').write_bytes(b'%PDF-1.4\nDeliberately damaged synthetic fixture, not a valid PDF.\n')
    from docx import Document
    doc=Document();doc.add_paragraph('SHIPPING INSTRUCTION');table=doc.add_table(rows=0,cols=2)
    for line in BASE.splitlines():
        label,value=line.split(': ',1);row=table.add_row();row.cells[0].text=label;row.cells[1].text=value
    doc.save(folder/'instruction-table.docx')
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Instruction';sheet.append(['SHIPPING INSTRUCTION'])
    for line in BASE.splitlines():sheet.append(line.split(': ',1))
    sheet.column_dimensions['A'].width=26;sheet.column_dimensions['B'].width=55;book.save(folder/'instruction-table.xlsx')
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image,ImageDraw,ImageFont
    pdf=Canvas(str(folder/'draft-native.pdf'))
    for n,line in enumerate(('DRAFT BILL OF LADING\n'+BASE).splitlines()):pdf.drawString(45,780-n*24,line)
    pdf.save()
    img=Image.new('RGB',(1500,1100),'white');draw=ImageDraw.Draw(img)
    try:font=ImageFont.truetype('DejaVuSans.ttf',30)
    except OSError:font=ImageFont.load_default(size=30)
    for n,line in enumerate(('DRAFT BILL OF LADING\n'+BASE).splitlines()):draw.text((50,70+n*90),line,fill='black',font=font)
    pdf=Canvas(str(folder/'draft-scanned.pdf'),pagesize=(750,550));pdf.drawImage(ImageReader(img),0,0,750,550);pdf.save()
    print(f'Synthetic fixtures written to {folder}')

if __name__=='__main__':main()
