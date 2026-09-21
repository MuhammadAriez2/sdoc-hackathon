"""
Format dispatcher: hand it a path, get back an `Extracted`.

The whole rest of the pipeline goes through this one function and never learns
what file type it was looking at. That is the seam that lets the four
extractors be built in parallel and swapped independently.

STATUS
------
    .txt    DONE
    .xlsx   TODO - track B
    .pdf    TODO - track C
    .docx   TODO - track C

The TODOs currently return `unreadable`, which makes those emails escalate to
NEEDS_REVIEW instead of producing a wrong answer. That is deliberate: an
honest "I cannot read this yet" scores better than a guess, and it means the
pipeline is correct end-to-end from the first run - you are improving
coverage, never fixing broken plumbing.
"""

import os

from .base import parse_labelled_text, unreadable
from ..contracts import Extracted


def extract(path: str, doc_role: str, root: str = ".") -> Extracted:
    """Read one attachment into the common `Extracted` shape.

    `path` is exactly the string from email["attachments"],
    e.g. "attachments/email_004_SI.txt".
    """
    full = os.path.join(root, path)
    ext = os.path.splitext(path)[1].lower()

    if not os.path.exists(full):
        return unreadable(doc_role, path, "file not found")
    if os.path.getsize(full) == 0:
        return unreadable(doc_role, path, "empty file")

    try:
        if ext == ".txt":
            return _extract_txt(full, doc_role, path)
        if ext == ".xlsx":
            return _extract_xlsx(full, doc_role, path)
        if ext == ".pdf":
            return _extract_pdf(full, doc_role, path)
        if ext == ".docx":
            return _extract_docx(full, doc_role, path)
    except Exception as exc:                      # noqa: BLE001
        # A parser blowing up is an unreadable document, not a crashed run.
        # 520 emails must always produce 520 outputs.
        return unreadable(doc_role, path, f"{type(exc).__name__}: {exc}")

    return unreadable(doc_role, path, f"unsupported format {ext}")


# ---------------------------------------------------------------------------
# .txt - working
# ---------------------------------------------------------------------------

def _extract_txt(full: str, doc_role: str, path: str) -> Extracted:
    with open(full, "r", encoding="utf-8", errors="replace") as fh:
        text = fh.read()
    if len(text.strip()) < 20:
        return unreadable(doc_role, path, "no readable text")
    return parse_labelled_text(text, doc_role, path)


# ---------------------------------------------------------------------------
# .xlsx - TODO (track B)
# ---------------------------------------------------------------------------

def _extract_xlsx(full: str, doc_role: str, path: str) -> Extracted:
    """TODO: read the workbook and join each row into 'Label: value' text.

    Sketch - this is genuinely most of it:

        import openpyxl
        wb = openpyxl.load_workbook(full, data_only=True)
        lines = []
        for ws in wb:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None]
                if len(cells) >= 2:
                    lines.append(f"{cells[0]}: {' '.join(cells[1:])}")
                elif cells:
                    lines.append(cells[0])
        return parse_labelled_text("\\n".join(lines), doc_role, path)

    Watch out for: weights arriving as bare numbers with no unit and no
    thousands separator (341715), and labels living in column A with the value
    spread across B and C. Add any new labels you meet to src/synonyms.py.
    """
    return unreadable(doc_role, path, "xlsx extractor not implemented yet")


# ---------------------------------------------------------------------------
# .pdf - TODO (track C)
# ---------------------------------------------------------------------------

def _extract_pdf(full: str, doc_role: str, path: str) -> Extracted:
    """TODO: pdfplumber, then the same parser.

        import pdfplumber
        with pdfplumber.open(full) as pdf:
            text = "\\n".join(p.extract_text() or "" for p in pdf.pages)
        if len(text.strip()) < 20:
            return unreadable(doc_role, path, "image-only PDF, no text layer")
        return parse_labelled_text(text, doc_role, path)

    Two things to know before you start:

    1. A PDF that yields 0 characters is one of the deliberate `unreadable`
       cases. Escalate it. Do NOT reach for OCR as the first move - test both
       and let the scorer decide.
    2. The layout is column-based, so a label and its value can end up on the
       same line separated by wide spaces rather than a colon. You may need a
       second pattern for "LABEL<lots of spaces>VALUE" alongside "Label: value".
    """
    return unreadable(doc_role, path, "pdf extractor not implemented yet")


# ---------------------------------------------------------------------------
# .docx - TODO (track C)
# ---------------------------------------------------------------------------

def _extract_docx(full: str, doc_role: str, path: str) -> Extracted:
    """TODO: python-docx. Every DOCX here puts the fields in a TABLE, so a
    paragraph-only reader will find nothing.

        import docx
        d = docx.Document(full)
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if len(cells) >= 2:
                    lines.append(f"{cells[0]}: {cells[1]}")
        return parse_labelled_text("\\n".join(lines), doc_role, path)

    The labels are bilingual - "PORT OF LOADING (装货港)". normalise_label()
    in src/synonyms.py already strips the bracket and the CJK characters, so
    they should map without extra work. Verify rather than assume.
    """
    return unreadable(doc_role, path, "docx extractor not implemented yet")
